from pathlib import Path

import pytest

from datamind.capabilities.ingest.formats import extract_document, extract_tabular


def test_extract_json_and_html_to_text(tmp_path: Path):
    payload = tmp_path / "payload.json"
    payload.write_text('{"name":"Ada","items":[1,2]}', encoding="utf-8")
    result = extract_document(payload)
    assert result.format == "json"
    assert '"name": "Ada"' in result.text

    page = tmp_path / "page.html"
    page.write_text("<h1>Report</h1><p>Total: 42</p>", encoding="utf-8")
    result = extract_document(page)
    assert result.text == "Report\nTotal: 42"


def test_extract_xlsx_returns_one_table_per_sheet(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Orders"
    first.append(["item", "cost"])
    first.append(["A", 10])
    second = workbook.create_sheet("Returns")
    second.append(["item", "reason"])
    second.append(["B", "damaged"])
    path = tmp_path / "workbook.xlsx"
    workbook.save(path)

    tables = extract_tabular(path)
    assert [table[0] for table in tables] == ["Orders", "Returns"]
    assert tables[0][2][0] == {"item": "A", "cost": 10}


def test_pdf_prefers_mineru_and_falls_back_to_pypdf(tmp_path: Path, monkeypatch):
    from datamind.capabilities.ingest import formats
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"placeholder")
    monkeypatch.delenv("DATAMIND_MINERU", raising=False)
    monkeypatch.setattr(formats, "_mineru_text", lambda path: ("# MinerU", [{"type": "markdown"}]))
    result = extract_document(pdf)
    assert result.text == "# MinerU"
    assert "parsed with MinerU API" in result.warnings

    monkeypatch.setattr(formats, "_mineru_text", lambda path: (_ for _ in ()).throw(RuntimeError("missing")))
    monkeypatch.setattr(formats, "_pdf_text", lambda path: ("fallback", [{"type": "page", "page": 1}]))
    result = extract_document(pdf)
    assert result.text == "fallback"
    assert any("fell back to pypdf" in warning for warning in result.warnings)


@pytest.mark.parametrize("dimension", ["A1:XFD1048576", "A1:A1"])
@pytest.mark.parametrize("prefixed", [False, True])
def test_xlsx_reads_actual_cells_despite_incorrect_declared_dimensions(tmp_path, dimension, prefixed):
    import hashlib
    import re
    import zipfile
    import openpyxl
    from datamind.config import DataSettings
    from datamind.workspace import Workspace
    from datamind.workspace_incremental import canonical_tables
    book = openpyxl.Workbook()
    book.active.append(["item", "amount"])
    book.active.append(["A", 12])
    path = tmp_path / "source.xlsx"; book.save(path); book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = re.sub(
        rb'<dimension ref="[^"]+"', f'<dimension ref="{dimension}"'.encode(),
        parts["xl/worksheets/sheet1.xml"])
    if prefixed:
        import xml.etree.ElementTree as ET
        ET.register_namespace("x", "http://schemas.openxmlformats.org/spreadsheetml/2006/main")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(ET.fromstring(parts["xl/worksheets/sheet1.xml"]))
    with zipfile.ZipFile(path, "w") as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    original = hashlib.sha256(path.read_bytes()).hexdigest()
    assert extract_tabular(path)[0][2] == [{"item": "A", "amount": 12}]
    workspace = Workspace(DataSettings(_env_file=None, data={"base_dir": tmp_path / "runtime"}))
    tables = canonical_tables(workspace, {"uri": str(path)}, path)
    assert tables[0]["columns"] == ["item", "amount", "_source_file", "_source_row"]
    from datamind.workspace_tables import table_rows
    assert next(table_rows(workspace.root, tables[0]))["_source_row"] == "2"
    assert tables[0]["rows"] == 1
    assert hashlib.sha256(path.read_bytes()).hexdigest() == original


def test_xlsx_title_row_does_not_hide_later_columns_and_empty_sheet_is_skipped(tmp_path):
    import openpyxl
    from datamind.config import DataSettings
    from datamind.workspace import Workspace
    from datamind.workspace_incremental import canonical_tables
    book = openpyxl.Workbook()
    book.active.append(["Quarterly report"])
    book.active.append(["sku", "amount"])
    book.active.append(["A", 12])
    book.create_sheet("Empty")
    path = tmp_path / "title.xlsx"; book.save(path); book.close()
    tables = extract_tabular(path)
    assert len(tables) == 1
    assert tables[0][1] == ["Quarterly report", "col_2"]
    assert tables[0][2][1]["col_2"] == 12
    workspace = Workspace(DataSettings(_env_file=None, data={"base_dir": tmp_path / "runtime"}))
    tables = canonical_tables(workspace, {"uri": str(path)}, path)
    assert len(tables) == 1 and tables[0]["columns"] == ["col_1", "col_2", "_source_file", "_source_row"]
    assert tables[0]["structure"]["header_row"] == 2
    assert tables[0]["rows"] == 2
