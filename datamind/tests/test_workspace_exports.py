"""Reusable evidence must match actual cells across updates and code execution."""
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from datamind.config import DataSettings
from datamind.core.errors import CapabilityError, ConfigError
from datamind.mcp import execute
from datamind.workspace import Workspace


@pytest.fixture
def ws(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("DATAMIND_DATA_ROOT", str(tmp_path / "runtime"))
    monkeypatch.setenv("DATAMIND__WORKSPACE__STORAGE_MODE", "snapshot")
    monkeypatch.setenv("DATAMIND__WORKSPACE__RETENTION", "all")
    monkeypatch.delenv("DATAMIND_PROFILE", raising=False)
    monkeypatch.delenv("DATAMIND_EXPORT_ROOT", raising=False)
    cfg = DataSettings(_env_file=None, data={"base_dir": tmp_path / "runtime"},
                       workspace={"storage_mode": "snapshot", "retention": "all", "auto_gc": False})
    source = tmp_path / "sources"
    source.mkdir()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "客户"
    for row in [("CRM summary",), (), ("指标", "学校", "合计"),
                ("客户数", 52, 214), ("续费率", "89%", "78%"),
                ("Note: do not silently invent units",)]:
        sheet.append(row)
    other = book.create_sheet("明细")
    other.append(["id", "amount"])
    other.append(["001", 10])
    other.append(["002", 20])
    book.save(source / "crm.xlsx")
    book.close()
    return Workspace(cfg), source


def mapping(source, **kwargs):
    path = source / "crm.xlsx"
    return [{"uri": str(path), "version": "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest(),
             "sheets": [{"sheet": "客户", "header_row": 3, "data_end_row": 5,
                         "fields": [{"column": "col_3", "name": "客户合计", "unit": "客户"}], **kwargs}]}]


def load_export(result):
    manifest = json.loads(Path(result["manifest"]["path"]).read_text())
    rows = [[json.loads(line) for line in Path(item["path"]).read_text().splitlines()]
            for item in manifest["items"] if item["kind"] == "table"]
    return manifest, rows


@pytest.mark.asyncio
async def test_batch_export_retains_cells_and_is_usable_without_original_parser(ws, monkeypatch):
    w, source = ws
    first = await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    from datamind import workspace_incremental
    monkeypatch.setattr(workspace_incremental, "canonical_tables", lambda *a, **k: pytest.fail("reparse"))
    # Missing originals cannot affect a published read or the generated analysis.
    (source / "crm.xlsx").unlink()
    async with w.reader(enable=set()) as reader:
        result = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])
        again = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])
        assert all(item["reused"] for item in again["items"])
        assert result["manifest"]["sha256"] == again["manifest"]["sha256"]
        manifest, rows = load_export(result)
        assert len(rows) == 2
        assert rows[0] == [{"指标": "客户数", "学校": "52", "客户合计": "214", "_source_row": 4},
                           {"指标": "续费率", "学校": "89%", "客户合计": "78%", "_source_row": 5}]
        assert manifest["workspace_revision"] == first["workspace_revision"]
        field = next(f for f in manifest["items"][0]["fields"] if f["key"] == "客户合计")
        assert field["source_column"] == "C" and field["unit"] == "客户"
        assert manifest["items"][0]["lineage"] == "source_row_and_column"
        program = "import json,sys; m=json.load(open(sys.argv[1])); p=next(x['path'] for x in m['items'] if x['sheet']=='明细'); print(sum(int(json.loads(l)['amount']) for l in open(p)))"
        child = subprocess.run([sys.executable, "-c", program, result["manifest"]["path"]], capture_output=True, text=True, check=True)
        assert child.stdout.strip() == "30"


@pytest.mark.asyncio
async def test_sql_asset_and_export_share_source_rows(ws):
    w, source = ws
    await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    async with w.reader() as reader:
        catalog = await reader.call("workspace_catalog", kind="tables", column="客户合计")
        assert catalog["sql_dialect"] == "sqlite"
        table = catalog["items"][0]
        assert table["structure"]["status"] == "confirmed"
        asset = await reader.call("workspace_read_asset", path="crm.xlsx", sheet="客户", offset=2, limit=1, columns=["col_3"])
        assert asset["rows"] == [["214"]] and asset["source_rows"] == [4]
        assert asset["source_columns"] == {"col_3": "C"}
        result = await reader.call("db_query_sql", sql=f'SELECT col_3, _source_row FROM "{table["table"]}" WHERE CAST(_source_row AS INTEGER)=4')
        assert result["rows"] == [["214", "4"]]
        raw = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx", "sheet": "客户"}], view_mode="raw")
        _, rows = load_export(raw)
        assert rows[0][-1]["col_1"].startswith("Note:")
        assert rows[0][1]["col_1"] == "指标"  # original row 3


@pytest.mark.asyncio
async def test_mapping_changes_reuse_parse_and_do_not_change_other_sources(ws, monkeypatch):
    w, source = ws
    (source / "other.csv").write_text("amount\n99\n")
    first = await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    from datamind import workspace_incremental
    monkeypatch.setattr(workspace_incremental, "canonical_tables", lambda *a, **k: pytest.fail("reparse"))
    changed = mapping(source, fields=[{"column": "col_3", "name": "总数"}])
    second = await w.update(str(source), surfaces=["db"], table_mappings=changed)
    a, b = [{s["relative_path"]: s for s in m["sources"]} for m in (first, second)]
    assert a["other.csv"]["checkpoint"] == b["other.csv"]["checkpoint"]
    assert second["metrics"]["parsed_table_files"] == 0
    assert a["crm.xlsx"]["checkpoint"] != b["crm.xlsx"]["checkpoint"]
    async with w.reader(first["workspace_revision"], enable=set()) as older:
        old_export = await older.call("workspace_export_bundle", requests=[{"path": "crm.xlsx", "sheet": "客户"}])
        third = await w.update(str(source), surfaces=["db"])
        assert third["sources"][0]["table_mapping"] == changed[0]["sheets"]
        async with w.reader(enable=set()) as newer:
            fresh = await newer.call("workspace_export_bundle", requests=[{"path": "crm.xlsx", "sheet": "客户"}])
        assert "客户合计" in load_export(old_export)[1][0][0]
        assert "总数" in load_export(fresh)[1][0][0]
        w.rollback(first["workspace_revision"], expected_revision=third["workspace_revision"])
        assert load_export(fresh)[0]["workspace_revision"] == third["workspace_revision"]
        assert "客户合计" in load_export(await older.call("workspace_export_bundle", requests=[{"path": "crm.xlsx", "sheet": "客户"}]))[1][0][0]


@pytest.mark.asyncio
async def test_changed_source_invalidates_mapping_and_retains_old_reader(ws):
    w, source = ws
    await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    async with w.reader(enable=set()) as older:
        stale = mapping(source)
        book = openpyxl.load_workbook(source / "crm.xlsx")
        book["客户"]["C4"] = 999
        book.save(source / "crm.xlsx")
        book.close()
        with pytest.raises(ConfigError, match="hash"):
            await w.update(str(source), surfaces=["db"], table_mappings=stale)
        second = await w.update(str(source), surfaces=["db"])
        assert second["sources"][0]["table_mapping"] == []
        result = await older.call("workspace_export_bundle", requests=[{"path": "crm.xlsx", "sheet": "客户"}])
        assert load_export(result)[1][0][0]["客户合计"] == "214"


@pytest.mark.asyncio
@pytest.mark.parametrize("change", [
    {"header_row": 99}, {"data_start_row": 2}, {"data_end_row": 99}, {"sheet": "absent"},
    {"fields": [{"column": "invented", "name": "fake"}]},
    {"fields": [{"column": "col_1", "name": " "}]},
    {"fields": [{"column": "col_1", "name": "a"}, {"column": "col_1", "name": "b"}]},
])
async def test_invalid_mapping_does_not_publish(ws, change):
    w, source = ws
    first = await w.update(str(source), surfaces=["db"])
    with pytest.raises(ConfigError):
        await w.update(str(source), surfaces=["db"], table_mappings=mapping(source, **change))
    assert w.current_revision == first["workspace_revision"]


@pytest.mark.asyncio
async def test_export_bounds_paths_cache_corruption_and_profile_scoping(ws, monkeypatch):
    w, source = ws
    await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    monkeypatch.setenv("DATAMIND_EXPORT_ROOT", str(source.parent / "shared-export"))
    async with w.reader(enable=set()) as reader:
        with pytest.raises(CapabilityError):
            await reader.call("workspace_export_bundle", requests=[{"path": "../private.csv"}])
        with pytest.raises(ValueError, match="max_rows"):
            await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}], max_rows=1)
        with pytest.raises(ValueError, match="Duplicate"):
            await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}, {"path": "crm.xlsx"}])
        exported = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])
        target = Path(exported["items"][0]["path"])
        assert "/default/exports/" in str(target)
        target.write_text("corrupt")
        repaired = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])
        assert not repaired["items"][0]["reused"]
        assert hashlib.sha256(target.read_bytes()).hexdigest() == repaired["items"][0]["sha256"]
        target.unlink()
        target.symlink_to(source / "crm.xlsx")
        with pytest.raises(CapabilityError, match="symlink"):
            await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])


@pytest.mark.asyncio
async def test_actual_mcp_batch_export_uses_operation_revision(ws):
    w, source = ws
    await w.update(str(source), surfaces=["db"], table_mappings=mapping(source))
    op = await execute("datamind_operation_begin", {"mode": "read"})
    try:
        result = await execute("datamind_workspace_export_bundle", {"operation_id": op["operation_id"], "requests": [{"path": "crm.xlsx"}]})
        assert result["workspace_revision"] == op["revision"]
        assert result["result"]["count"] == 2
        assert len(json.dumps(result)) < 8000
    finally:
        await execute("datamind_operation_close", {"operation_id": op["operation_id"]})


@pytest.mark.asyncio
async def test_mcp_publish_accepts_hashed_table_mappings_and_checks_receipt(ws):
    w, source = ws
    writer = await execute("datamind_operation_begin", {"mode": "write"})
    try:
        result = await execute("datamind_workspace_publish", {"operation_id": writer["operation_id"],
            "path": str(source), "surfaces": ["db"], "table_mappings": mapping(source)})
        assert all(r["status"] == "stored" for r in result["result"]["results"])
        async with w.reader(enable=set()) as reader:
            catalog = await reader.call("workspace_catalog", kind="tables", column="客户合计")
            assert catalog["items"][0]["structure"]["header_row"] == 3
    finally:
        await execute("datamind_operation_close", {"operation_id": writer["operation_id"]})


@pytest.mark.asyncio
async def test_snapshot_gc_keeps_shared_canonical_tables(ws):
    w, source = ws
    await w.update(str(source), surfaces=["db"])
    w.configure_retention("latest")
    await w.update(str(source), surfaces=["db"])
    await w.collect(dry_run=False)
    async with w.reader(enable=set()) as reader:
        result = await reader.call("workspace_export_bundle", requests=[{"path": "crm.xlsx"}])
        assert result["count"] == 2


@pytest.mark.asyncio
async def test_text_exports_reuse_frozen_parse_without_opening_backends(ws, monkeypatch):
    w, source = ws
    note = source / "notes.md"
    note.write_text("第一版来源\nRevenue is reported in USD.")
    first = await w.update(str(source), surfaces=["db"])
    note.write_text("Second version")
    await w.update(str(source), surfaces=["db"])
    note.unlink()
    from datamind.capabilities.ingest import formats
    monkeypatch.setattr(formats, "extract_document", lambda *a: pytest.fail("unexpected parsing"))
    from datamind import runtime
    monkeypatch.setattr(runtime, "build_db_service", lambda *a, **k: pytest.fail("SQL backend opened"))
    monkeypatch.setattr(runtime, "build_embedding", lambda *a, **k: pytest.fail("embedding opened"))
    async with w.reader(first["workspace_revision"], enable=set()) as reader:
        result = await reader.call("workspace_export_bundle", requests=[{"path": "notes.md"}, {"path": "crm.xlsx"}])
        manifest, _ = load_export(result)
        text = next(item for item in manifest["items"] if item["kind"] == "text")
        assert Path(text["path"]).read_text() == "第一版来源\nRevenue is reported in USD."
        assert text["parsed_cache_reused"] and result["count"] == 3


@pytest.mark.asyncio
async def test_byte_limit_fails_instead_of_returning_truncated_text(ws):
    w, source = ws
    (source / "large.md").write_text("原文" * 2000)
    await w.update(str(source), surfaces=["db"])
    async with w.reader(enable=set()) as reader:
        with pytest.raises(ValueError, match="max_bytes"):
            await reader.call("workspace_export_bundle", requests=[{"path": "large.md"}], max_bytes=1024)


@pytest.mark.asyncio
async def test_export_profile_directory_cannot_redirect_to_other_data(ws, monkeypatch):
    w, source = ws
    await w.update(str(source), surfaces=["db"])
    cache = source.parent / "exports"
    cache.mkdir()
    (cache / "default").symlink_to(source, target_is_directory=True)
    monkeypatch.setenv("DATAMIND_EXPORT_ROOT", str(cache))
    with pytest.raises(CapabilityError, match="symlink"):
        async with w.reader(enable=set()):
            pytest.fail("unsafe export root accepted")


def test_legacy_export_does_not_fabricate_original_cell_positions(tmp_path):
    from datamind.workspace_discovery import WorkspaceView
    from datamind.workspace_exports import export_bundle
    table = {"name": "legacy", "sheet": "Sheet1", "columns": ["label", "_source_file"],
             "headers": ["label"], "rows": 1, "records": [{"label": "retained", "_source_file": "/old.csv"}]}
    (tmp_path / "artifact.json").write_text(json.dumps({"tables": [table]}))
    manifest = {"workspace_revision": "legacy-revision", "sources": [{"uri": "/old.csv",
        "relative_path": "old.csv", "version": "sha256:fixture", "artifact": "artifact.json"}]}
    view = WorkspaceView(manifest, tmp_path)
    result = export_bundle(view, [{"path": "old.csv"}], tmp_path / "exports")
    exported, rows = load_export(result)
    assert exported["items"][0]["lineage"] == "unavailable_legacy_republish_required"
    assert rows[0][0] == {"label": "retained", "_source_row": None}
    asset = view.read_asset(path="old.csv")
    assert asset["source_rows"] is None and asset["source_columns"] is None
